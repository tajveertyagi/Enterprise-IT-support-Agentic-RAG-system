from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader,TextLoader
from docx import Document as DocxDocument


SUPPORTED={".pdf", ".txt", ".docx",".md"}


def load_file(path:Path) ->list[Document]:
    suffix=path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {suffix}")
    
    if suffix==".pdf":
        loader=PyPDFLoader(str(path))
        return loader.load()
    elif suffix in {".txt",".md"}:
        loader=TextLoader(str(path),encoding="utf-8")
        return loader.load()
    elif suffix==".docx":
        doc=DocxDocument(str(path))
        text="\n".join([para.text for para in doc.paragraphs])
        return [Document(page_content=text, metadata={"source": str(path)})]
   
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def split_documents(docs:Iterable[Document],chunk_size:int=900,chunk_overlap:int=150,add_start_index=True) ->list[Document]:
    splitter=RecursiveCharacterTextSplitter(chunk_size=chunk_size,chunk_overlap=chunk_overlap)
    return splitter.split_documents(list(docs))
